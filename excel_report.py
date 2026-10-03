"""Export massal hasil kuesioner ke Excel (.xlsx).

Sheet:
  1. Rekap           - satu baris per responden, skor A-E, total, rata-rata, level (rumus)
  2. Detail Jawaban  - satu baris per jawaban: teks jawaban, skor, alasan, sumber
  3. Keterangan      - info export, kode kompetensi, skala skor
"""

from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

from questions import KOMPETENSI, KOMPETENSI_BY_KODE, LABEL_SKOR

FONT = "Arial"
WARNA_UTAMA = "1F3A5F"
WARNA_LATAR = "F2F5F9"
GARIS = Side(style="thin", color="C9D3DF")
BORDER = Border(left=GARIS, right=GARIS, top=GARIS, bottom=GARIS)
KODE = [k["kode"] for k in KOMPETENSI]

# Rumus level dari nilai rata-rata/skor (1-5)
RUMUS_LEVEL = ('=IF({ref}="","",CHOOSE(ROUND({ref},0),'
               '"Sangat kurang","Kurang","Memadai","Kuat","Sangat kuat"))')


def _header(ws, judul: list[str], baris: int = 1) -> None:
    for col, teks in enumerate(judul, start=1):
        c = ws.cell(row=baris, column=col, value=teks)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=WARNA_UTAMA)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER


def _gaya_sel(c, tengah: bool = False, wrap: bool = False, tebal: bool = False) -> None:
    c.font = Font(name=FONT, bold=tebal)
    c.border = BORDER
    c.alignment = Alignment(horizontal="center" if tengah else "left",
                            vertical="top", wrap_text=wrap)


def _lebar(ws, lebar: list[float]) -> None:
    for i, w in enumerate(lebar, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _sheet_rekap(ws, respondents: list[dict]) -> None:
    judul = (["No", "Nama", "Posisi", "Unit Kerja", "Lama Bekerja (th)", "Tanggal Pengisian"]
             + [f"{k['kode']}. {k['nama']}" for k in KOMPETENSI]
             + ["Total (maks 25)", "Rata-rata", "Level", "Jumlah Dinilai Aturan"])
    _header(ws, judul)
    ws.row_dimensions[1].height = 48

    kol_a = 7                                   # kolom skor A
    kol_e = kol_a + len(KODE) - 1               # kolom skor E
    huruf_a, huruf_e = get_column_letter(kol_a), get_column_letter(kol_e)
    kol_total, kol_rata, kol_level, kol_aturan = kol_e + 1, kol_e + 2, kol_e + 3, kol_e + 4
    huruf_rata = get_column_letter(kol_rata)

    for i, r in enumerate(respondents, start=1):
        baris = i + 1
        skor = {j["kode_kompetensi"]: j["skor"] for j in r["jawaban"]}
        aturan = sum(1 for j in r["jawaban"] if j["sumber"] == "Aturan")
        nilai = [i, r["nama"], r["posisi"], r["unit_kerja"], float(r["lama_bekerja"]),
                 r["tanggal_pengisian"]] + [skor.get(k) for k in KODE]
        for col, v in enumerate(nilai, start=1):
            c = ws.cell(row=baris, column=col, value=v)
            _gaya_sel(c, tengah=(col == 1 or col >= 5))
        ws.cell(row=baris, column=6).number_format = "DD-MM-YYYY"
        ws.cell(row=baris, column=5).number_format = "0.0"

        rentang = f"{huruf_a}{baris}:{huruf_e}{baris}"
        rumus = {
            kol_total: f"=IF(COUNT({rentang})=0,\"\",SUM({rentang}))",
            kol_rata: f"=IFERROR(AVERAGE({rentang}),\"\")",
            kol_level: RUMUS_LEVEL.format(ref=f"{huruf_rata}{baris}"),
        }
        for col, f in rumus.items():
            c = ws.cell(row=baris, column=col, value=f)
            _gaya_sel(c, tengah=True)
        ws.cell(row=baris, column=kol_rata).number_format = "0.00"
        c = ws.cell(row=baris, column=kol_aturan, value=aturan)
        _gaya_sel(c, tengah=True)

    # Baris rata-rata: SUBTOTAL(101,...) = AVERAGE yang mengikuti filter (baris tersembunyi diabaikan)
    akhir = len(respondents) + 1
    baris_rata = akhir + 1
    c = ws.cell(row=baris_rata, column=2, value="Rata-rata (mengikuti filter)")
    for col in range(1, kol_aturan + 1):
        sel = ws.cell(row=baris_rata, column=col)
        sel.fill = PatternFill("solid", fgColor=WARNA_LATAR)
        _gaya_sel(sel, tengah=(col >= 5), tebal=True)
    for col in list(range(kol_a, kol_e + 1)) + [kol_total, kol_rata]:
        huruf = get_column_letter(col)
        sel = ws.cell(row=baris_rata, column=col,
                      value=f"=IFERROR(SUBTOTAL(101,{huruf}2:{huruf}{akhir}),\"\")")
        _gaya_sel(sel, tengah=True, tebal=True)
        sel.fill = PatternFill("solid", fgColor=WARNA_LATAR)
        sel.number_format = "0.00"
    sel = ws.cell(row=baris_rata, column=kol_level,
                  value=RUMUS_LEVEL.format(ref=f"{huruf_rata}{baris_rata}"))
    _gaya_sel(sel, tengah=True, tebal=True)
    sel.fill = PatternFill("solid", fgColor=WARNA_LATAR)

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(kol_aturan)}{akhir}"
    _lebar(ws, [5, 26, 24, 18, 11, 13] + [14] * len(KODE) + [11, 10, 14, 13])


def _sheet_detail(ws, respondents: list[dict]) -> None:
    judul = ["ID Responden", "Nama", "Posisi", "Unit Kerja", "Tanggal Pengisian",
             "Kode", "Kompetensi", "Jawaban", "Skor", "Level", "Alasan Penilaian",
             "Sumber", "Model AI"]
    _header(ws, judul)
    ws.row_dimensions[1].height = 32

    baris = 2
    for r in respondents:
        for j in r["jawaban"]:
            k = KOMPETENSI_BY_KODE[j["kode_kompetensi"]]
            nilai = [r["id"], r["nama"], r["posisi"], r["unit_kerja"], r["tanggal_pengisian"],
                     k["kode"], k["nama"], j["jawaban"], j["skor"], None, j["alasan"],
                     j["sumber"], j.get("model") or "-"]
            for col, v in enumerate(nilai, start=1):
                c = ws.cell(row=baris, column=col, value=v)
                _gaya_sel(c, tengah=col in (1, 5, 6, 9, 10, 12), wrap=col in (7, 8, 11))
            ws.cell(row=baris, column=5).number_format = "DD-MM-YYYY"
            c = ws.cell(row=baris, column=10, value=RUMUS_LEVEL.format(ref=f"I{baris}"))
            _gaya_sel(c, tengah=True)
            baris += 1

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:M{max(baris - 1, 1)}"
    _lebar(ws, [11, 24, 22, 16, 13, 7, 24, 70, 7, 13, 50, 9, 20])


def _sheet_keterangan(ws, jumlah: int, keterangan: str) -> None:
    ws["A1"] = "Rekap Hasil Kuesioner Kompetensi"
    ws["A1"].font = Font(name=FONT, bold=True, size=14, color=WARNA_UTAMA)
    info = [
        ("Dicetak", datetime.now().strftime("%d-%m-%Y %H:%M")),
        ("Jumlah responden", jumlah),
        ("Filter", keterangan or "Semua responden"),
    ]
    baris = 3
    for label, nilai in info:
        ws.cell(row=baris, column=1, value=label).font = Font(name=FONT, bold=True)
        ws.cell(row=baris, column=2, value=nilai).font = Font(name=FONT)
        baris += 1

    baris += 1
    ws.cell(row=baris, column=1, value="Kode Kompetensi").font = Font(name=FONT, bold=True)
    baris += 1
    for k in KOMPETENSI:
        ws.cell(row=baris, column=1, value=k["kode"]).font = Font(name=FONT)
        ws.cell(row=baris, column=2, value=k["nama"]).font = Font(name=FONT)
        baris += 1

    baris += 1
    ws.cell(row=baris, column=1, value="Skala Skor").font = Font(name=FONT, bold=True)
    baris += 1
    for skor in sorted(LABEL_SKOR, reverse=True):
        ws.cell(row=baris, column=1, value=skor).font = Font(name=FONT)
        ws.cell(row=baris, column=2, value=LABEL_SKOR[skor]).font = Font(name=FONT)
        baris += 1

    baris += 1
    catatan = [
        "Sumber 'AI' = dinilai Gemini; 'Aturan' = dinilai aturan kata kunci tanpa AI (fallback).",
        "Kolom Total, Rata-rata, dan Level di sheet Rekap dihitung dengan rumus Excel.",
        "Baris 'Rata-rata (mengikuti filter)' otomatis menyesuaikan saat filter kolom dipakai.",
        "Data ini bersifat rahasia.",
    ]
    for teks in catatan:
        ws.cell(row=baris, column=1, value=teks).font = Font(name=FONT, italic=True,
                                                            color="555555")
        baris += 1
    _lebar(ws, [20, 60])


def build_excel(respondents: list[dict], keterangan: str = "") -> bytes:
    wb = Workbook()
    ws_rekap = wb.active
    ws_rekap.title = "Rekap"
    _sheet_rekap(ws_rekap, respondents)
    _sheet_detail(wb.create_sheet("Detail Jawaban"), respondents)
    _sheet_keterangan(wb.create_sheet("Keterangan"), len(respondents), keterangan)

    # Pastikan Excel menghitung ulang semua rumus saat file dibuka
    wb.calculation = CalcProperties(fullCalcOnLoad=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
